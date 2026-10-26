"""Workbook previews, explicit Pending column mapping and three-way comparison."""
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook
from domain import FIELDS, normalize_item, payload, ValidationError
from excel_sync import MAX_FILE, MAX_UNCOMPRESSED


def workbook_view(content, config):
    """Preview visible sheets; map one explicitly configured sheet to Pending."""
    if len(content) > MAX_FILE:
        raise ValidationError('Excel은 20MB 이하여야 합니다.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_UNCOMPRESSED:
                raise ValidationError('압축 해제 크기가 40MB를 초과합니다.')
        book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, OSError, KeyError):
        raise ValidationError('정상적인 .xlsx 파일이 아닙니다.') from None
    try:
        previews = {}
        budget = 100000
        for sheet in book:
            if sheet.sheet_state != 'visible':
                continue
            limit = min(500, max(0, budget // 50))
            if not limit:
                break
            values = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row or 1, limit), max_col=min(sheet.max_column or 1, 50), values_only=True))
            budget -= sum(len(r) for r in values)
            previews[sheet.title] = {'rows': values, 'truncated': (sheet.max_row or 0) > limit or (sheet.max_column or 0) > 50}
        name = config.get('pending_sheet', 'PENDING')
        if name not in book.sheetnames:
            return previews, None, f'PENDING 연결 대상 시트({name})가 없습니다. 표 조회만 가능합니다.'
        sheet = book[name]
        header_row = int(config.get('header_row', 1))
        if not 1 <= header_row <= 100:
            raise ValidationError('머리글 행은 1~100이어야 합니다.')
        if (sheet.max_column or 0) > 200:
            raise ValidationError('PENDING 연결 시트는 200열 이하여야 합니다.')
        if (sheet.max_row or 0) > header_row + 5000:
            raise ValidationError('PENDING 연결 시트는 5,000행 이하여야 합니다.')
        header = next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
        names = [str(v).strip() if v is not None else '' for v in header]
        columns = config.get('columns', {})
        indexes = {}
        for field in FIELDS:
            source = columns.get(field, field)
            if names.count(source) != 1:
                return previews, None, f'열 연결 확인 필요: {field} → {source} (열이 없거나 중복됩니다).'
            indexes[field] = names.index(source)
        result = []; seen = set()
        state_map = config.get('status_values', {})
        category_map = config.get('category_values', {})
        for number, values in enumerate(sheet.iter_rows(min_row=header_row+1, values_only=True), header_row+1):
            if all(v is None for v in values):
                continue
            raw = {field: values[i] if i < len(values) else None for field, i in indexes.items()}
            raw['상태'] = state_map.get(str(raw['상태']).strip(), raw['상태'])
            raw['분야'] = category_map.get(str(raw['분야']).strip(), raw['분야'])
            try:
                row = payload(normalize_item(raw))
            except ValidationError as exc:
                return previews, None, f'{name} {number}행: {exc}'
            if row['ID'] in seen:
                return previews, None, f'{name} {number}행: ID가 중복됩니다.'
            seen.add(row['ID']); result.append(row)
        # Required Pending values must be literal, never silently use stale formula caches.
        formula_book = load_workbook(BytesIO(content), read_only=True, data_only=False)
        try:
            for cells in formula_book[name].iter_rows(min_row=header_row+1):
                if any(i < len(cells) and cells[i].data_type == 'f' for i in indexes.values()):
                    return previews, None, 'PENDING 연결 열에는 수식 대신 값을 입력해주세요.'
        finally:
            formula_book.close()
        return previews, result, None
    finally:
        book.close()


def source_plan(current, previous, incoming):
    """Three-way comparison: remote-only edits apply; divergent edits conflict."""
    active = {r['ID']: r for r in current}
    baseline = {r['ID']: r for r in previous}
    plan = {'changes': [], 'conflicts': [], 'missing': []}
    for row in incoming:
        identifier = row['ID']; old = active.get(identifier); base = baseline.get(identifier)
        if old and payload(old) == row:
            continue
        if base == row:
            continue  # only the app changed: retain the app edit
        if old and (base is None or payload(old) != base):
            plan['conflicts'].append({'ID': identifier, '사유': 'Excel과 앱의 값이 다릅니다. 양쪽 값을 일치시킨 뒤 다시 동기화해주세요.'})
        elif old is None and base is not None:
            plan['conflicts'].append({'ID': identifier, '사유': '기존 동기화 항목이 DB에 없습니다. 관리자 확인이 필요합니다.'})
        else:
            plan['changes'].append({'id': identifier, 'after': row, 'expected_version': old['버전'] if old else 0})
    plan['missing'] = [r for r in previous if r['ID'] not in {v['ID'] for v in incoming}]
    return plan
