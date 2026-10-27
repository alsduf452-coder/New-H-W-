"""Parser for the OCTA H/W workbook layout. All values retain their sign and position."""
from datetime import date, datetime
from io import BytesIO
import math
import re
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries, get_column_letter
from domain import ValidationError

SCHEMA_VERSION = 'octa-hw-1'
REQUIRED = ('이력', 'Spacing', 'Diff Flatness', 'Susc Flatness')
CHAMBERS = tuple('ABCDE')


def scalar(value):
    if isinstance(value, datetime):
        return value.isoformat(sep=' ')
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool) and math.isfinite(value)


def measured_date(value, location):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    try:
        return date.fromisoformat(str(value).strip()).isoformat()
    except ValueError:
        raise ValidationError(f'{location}: 측정/변경 일자가 없거나 날짜 형식이 다릅니다.') from None


def parse_hw(content):
    book = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        if any(name not in book.sheetnames for name in REQUIRED):
            raise ValidationError('H/W 양식에는 이력, Spacing, Diff Flatness, Susc Flatness 시트가 필요합니다.')
        data = {}
        for name in REQUIRED:
            sheet = book[name]
            if (sheet.max_row or 0) > 5000 or (sheet.max_column or 0) > 200:
                raise ValidationError(f'{name}: H/W 해석 범위(5,000행·200열)를 초과합니다.')
            data[name] = list(sheet.iter_rows(values_only=True))
        def cell(name, row, col):
            rows = data[name]
            return scalar(rows[row-1][col-1]) if 0 < row <= len(rows) and 0 < col <= len(rows[row-1]) else None
        def raw(name, row, col):
            rows = data[name]
            return rows[row-1][col-1] if 0 < row <= len(rows) and 0 < col <= len(rows[row-1]) else None
        header = [cell('이력', 3, c) for c in range(2, 13)]
        expected = ['Chamber', '일자', '변경이력', 'Diffuser', 'Backing Plate', 'Susceptor', 'Shadow Frame', 'RFG', 'Matcher', 'Reactor', 'Generator']
        if header != expected:
            raise ValidationError('이력 시트 3행의 열 구조가 변경됐습니다. 자동 반영을 보류합니다.')
        # Fill only actual merged reason cells; do not forward-fill arbitrary empty rows.
        merged_reasons = {}
        with ZipFile(BytesIO(content)) as archive:
            path = book['이력']._worksheet_path.lstrip('/')
            root = ET.fromstring(archive.read(path))
            for node in root.findall('.//{*}mergeCell'):
                left, top, right, bottom = range_boundaries(node.attrib['ref'])
                if left == right == 4:
                    for row in range(top, bottom+1):
                        merged_reasons[row] = cell('이력', top, 4)
        parts = expected[3:]
        parts = ['RPSC Reactor' if p == 'Reactor' else p for p in parts]
        history = []
        for r in range(4, len(data['이력'])+1):
            values = [cell('이력', r, c) for c in range(2, 13)]
            if all(v is None for v in values):
                continue
            chamber = str(values[0] or '').strip()
            if chamber not in CHAMBERS:
                raise ValidationError(f'이력 B{r}: Chamber A~E 값이 필요합니다.')
            when = measured_date(raw('이력', r, 3), f'이력 C{r}')
            serials = {part: str(cell('이력', r, c)).strip() for part, c in zip(parts, range(5, 13)) if cell('이력', r, c) not in (None, '')}
            history.append({'chamber': chamber, 'date': when, 'reason': merged_reasons.get(r, cell('이력', r, 4)) or '', 'parts': serials, 'source': f'이력!B{r}:L{r}', 'source_row': r})
        measurements = []; warnings = []
        for sheet_name, kind in [('Spacing', 'Spacing'), ('Diff Flatness', 'Diffuser Flatness'), ('Susc Flatness', 'Susceptor Flatness')]:
            is_spacing = kind == 'Spacing'
            anchors = {2 + 9*i if is_spacing else 3 + 9*i: chamber for i, chamber in enumerate(CHAMBERS)}
            # Verify chamber bands instead of silently assigning data to a changed layout.
            header_row = 3 if is_spacing else 4
            for c, chamber in anchors.items():
                candidates = [cell(sheet_name, header_row, k) for k in range(max(1, c-1), c+1)]
                if not any(re.sub(r'[^A-Z]', '', str(v).upper()) == 'PC'+chamber for v in candidates):
                    raise ValidationError(f'{sheet_name}: Chamber {chamber}의 위치가 변경됐습니다.')
            for r, values in enumerate(data[sheet_name], 1):
                for c, value in enumerate(values, 1):
                    if str(value).strip().lower() != 'date':
                        continue
                    if c not in anchors:
                        raise ValidationError(f'{sheet_name}!{get_column_letter(c)}{r}: 지원하지 않는 측정 위치입니다.')
                    chamber = anchors[c]
                    when = measured_date(raw(sheet_name, r, c+1), f'{sheet_name}!{get_column_letter(c+1)}{r}')
                    if cell(sheet_name, r+1, c) != 'Diff. S/N' or cell(sheet_name, r+1, c+3) != 'Susc. S/N':
                        raise ValidationError(f'{sheet_name} {r}행: S/N 배치가 변경됐습니다.')
                    if is_spacing:
                        labels = [['①','⑥','②'], ['⑨','⑤','⑦'], ['④','⑧','③']]
                        for i in range(3):
                            for j in range(3):
                                if cell(sheet_name, r+3+2*i, c+1+2*j) != labels[i][j]:
                                    raise ValidationError(f'{sheet_name} {r}행: 9개 측정점 배치가 변경됐습니다.')
                        coords = [[(r+4+2*i, c+1+2*j) for j in range(3)] for i in range(3)]
                        reason = cell(sheet_name, r-1, c)
                    else:
                        coords = [[(r+3+i, c+j) for j in range(7)] for i in range(5)]
                        reason = cell(sheet_name, r+2, c)
                    grid = [[cell(sheet_name, row, col) for row,col in line] for line in coords]
                    numbers = [v for line in grid for v in line if numeric(v)]
                    center = grid[1][1] if is_spacing else grid[2][3]
                    record = {'id': f'{sheet_name}:{get_column_letter(c)}{r}', 'kind': kind, 'chamber': chamber, 'date': when, 'reason': reason or '', 'diffuser_sn': cell(sheet_name,r+1,c+1), 'susceptor_sn': cell(sheet_name,r+1,c+4), 'grid': grid, 'cells': [[f'{get_column_letter(col)}{row}' for row,col in line] for line in coords], 'source_sheet': sheet_name, 'source_row': r, 'center': center, 'count': len(numbers), 'total': sum(len(line) for line in grid), 'min': min(numbers) if numbers else None, 'max': max(numbers) if numbers else None, 'range': max(numbers)-min(numbers) if numbers else None}
                    if any(v not in (None, '') and not numeric(v) for line in grid for v in line):
                        warnings.append(f'{record["id"]}: 숫자가 아닌 측정값은 원문 표시하고 통계에서 제외했습니다.')
                    measurements.append(record)
        if not history or any(not any(m['kind']==kind for m in measurements) for kind in ('Spacing','Diffuser Flatness','Susceptor Flatness')):
            raise ValidationError('H/W 이력 또는 측정 기록이 없습니다. 기존 데이터를 유지합니다.')
        installed = {chamber: {} for chamber in CHAMBERS}
        for event in sorted(history, key=lambda v:(v['date'],v['source_row'])):
            for part,sn in event['parts'].items():
                installed[event['chamber']][part] = {'sn':sn,'date':event['date'],'source':event['source']}
        return {'schema': SCHEMA_VERSION, 'history': history, 'measurements': measurements, 'installed': installed, 'parts': parts, 'warnings': warnings}
    finally:
        book.close()
