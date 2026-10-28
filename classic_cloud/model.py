from datetime import date, datetime
from io import BytesIO
import re
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pending_octa import target_date

MAIN = ['HW(설계)','공정','SW','전장','기타']
COLS = ['ID','범위','설비','분야','상태','Unit','내용','담당자','목표일','발생일','원인','진행 사항','조치내용','완료일','조치자','45호기 적용 여부','시트']


def text(value):
    return '' if value is None else str(value).strip()


def pending_frame(snapshot):
    previews = snapshot.get('previews',{})
    details = previews.get('__pending_data__',{}).get('details',[])
    rows=[]
    for detail in details:
        rows.append({'ID':detail['ID'],'범위':'주요 5개 분야','설비':text(detail.get('호기')) or '미기재','분야':detail.get('시트',''),'상태':'완료' if text(detail.get('완료여부')).upper() in ('O','○','완료') else '미완료','Unit':text(detail.get('Unit명')),'내용':text(detail.get('불합리 및 이슈 내용')),'담당자':text(detail.get('담당자')),'목표일':text(detail.get('목표일 원문')),'발생일':text(detail.get('발생일')),'원인':text(detail.get('사유 or 원인')),'진행 사항':text(detail.get('진행 사항',detail.get('진행 사항 (미팅내용)'))),'조치내용':text(detail.get('조치내용')),'완료일':text(detail.get('완료일')),'조치자':text(detail.get('담당자(조치자)')),'45호기 적용 여부':text(detail.get('45호기 적용 여부')),'시트':detail.get('시트','')})
    for name in ['고객사요청','software불합리','설비알람']:
        for n,row in enumerate(previews.get(name,{}).get('rows',[]),1):
            val=lambda col: text(row[col-1]) if col<=len(row) else ''
            number=val(2 if name=='고객사요청' else 1)
            if not re.fullmatch(r'\d+(?:\.0)?',number):continue
            item={key:'' for key in COLS}
            item.update(ID=f'aux-{name}-{n}',범위='보조 시트',분야=name,시트=name,상태='미완료')
            if name=='고객사요청':
                item.update(Unit=val(4),내용=val(5),원인=val(6) if n<39 else '',**{'진행 사항':val(11),'45호기 적용 여부':val(13)})
                if n<39:
                    item.update(설비='미기재',상태='완료' if all(val(c) in ('O','○','완료') for c in (8,9)) else '미완료')
                else:
                    item.update(설비=val(7),발생일=val(6),목표일=val(8),완료일=val(9),상태='완료' if val(10) in ('O','○','완료') else '미완료')
            elif name=='software불합리':
                item.update(설비=val(2),상태='완료' if val(3) in ('O','○','완료') else '미완료',목표일=val(4),완료일=val(5),내용=val(6),**{'진행 사항':val(7)})
            else:
                item.update(설비=val(3),발생일=val(2),Unit=val(4),내용=val(5)+'\n'+val(6),원인=val(7),조치내용=val(8),**{'진행 사항':val(9)})
            item['설비']=item['설비'] or '미기재'
            if item['내용']:rows.append(item)
    return pd.DataFrame(rows,columns=COLS)


def filter_pending(frame, scope='주요 5개 분야', equipment=(), categories=(), states=(), query=''):
    result=frame.copy()
    if scope!='전체 시트':result=result[result['범위']==scope]
    for field,values in [('설비',equipment),('분야',categories),('상태',states)]:
        if values:result=result[result[field].isin(values)]
    for term in query.split():
        if result.empty:break
        result=result[result.astype(str).agg(' '.join,axis=1).str.contains(term,case=False,regex=False)]
    return result


def summary(frame, group):
    if frame.empty:return pd.DataFrame(columns=[group,'완료','미완료','전체','완료율(%)'])
    grouped=pd.crosstab(frame[group],frame['상태']).reindex(columns=['완료','미완료'],fill_value=0)
    grouped['전체']=grouped.sum(axis=1);grouped['완료율(%)']=(grouped['완료']/grouped['전체']*100).round(1)
    return grouped.reset_index()


def deadlines(frame, today):
    counts={'기한초과':0,'임박':0,'여유':0};unresolved=0
    for row in frame.to_dict('records'):
        if row['상태']!='미완료':continue
        value=row['목표일']
        # Preview date cells are ISO datetime strings; retain raw text elsewhere.
        if re.fullmatch(r'\d{4}-\d{2}-\d{2}[ T]00:00:00',value):value=value[:10]
        normalized,_=target_date(value)
        if not normalized:unresolved+=1;continue
        days=(date.fromisoformat(normalized)-today).days
        counts['기한초과' if days<0 else '임박' if days<=7 else '여유']+=1
    return counts,unresolved


def revision(snapshot):
    previews=snapshot.get('previews',{})
    name=next((n for n in previews if re.sub(r'\s','',n).lower()=='rev관리'),None)
    if not name:return {'history':[],'summary':[],'title':'','line':'','process':''}
    rows=previews[name]['rows']
    def val(r,c):return text(rows[r-1][c-1]) if r<=len(rows) and c<=len(rows[r-1]) else ''
    totals=[]
    for c in (2,4,6,8,10,12,14):
        if not val(5,c):continue
        rate=val(8,c+1)
        try:rate=f'{float(rate)*100:.1f}%'
        except ValueError:pass
        totals.append({'구분':val(5,c),'총 건수':val(7,c),'조치 건수':val(7,c+1),'조치율':rate or '—'})
    fields=['REV','구분','주요 변경내용','변경일자','작성자','변경점 Key Word / 비고']
    history=[dict(zip(fields,[val(r,c) for c in (2,3,4,10,11,12)])) for r in range(11,len(rows)+1) if re.fullmatch(r'\d+(?:\.0)?',val(r,2)) and val(r,4)]
    return {'title':val(2,4),'line':val(2,13),'process':val(3,13),'summary':totals,'history':history}


def workbook_export(tables):
    book=Workbook();book.remove(book.active)
    for name,frame in tables.items():
        sheet=book.create_sheet(name[:31]);sheet.append([str(c) for c in frame.columns])
        for row in frame.itertuples(index=False,name=None):
            sheet.append([None if pd.isna(v) else v for v in row])
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str):cell.data_type='s'
                cell.alignment=Alignment(vertical='top',wrap_text=True)
        for cell in sheet[1]:cell.font=Font(bold=True);cell.fill=PatternFill('solid',fgColor='CBD5E1')
        for i,column in enumerate(frame.columns,1):
            letter=sheet.cell(1,i).column_letter
            sheet.column_dimensions[letter].width=65 if str(column) in ('내용','주요 변경내용','변경점 Key Word / 비고') else 22
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        sheet.sheet_view.zoomScale=85
        sheet.sheet_properties.pageSetUpPr.fitToPage=True
        sheet.page_setup.orientation='landscape';sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    output=BytesIO();book.save(output);return output.getvalue()
