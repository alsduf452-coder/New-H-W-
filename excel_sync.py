"""고정된 PENDING 연동 양식. 수식·중복 ID·잘못된 날짜를 묵인하지 않습니다."""
from io import BytesIO
from zipfile import ZipFile, BadZipFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from domain import FIELDS, FIELDS_WITH_VERSION, ValidationError, normalize_item

MAX_FILE = 20 * 1024 * 1024
MAX_UNCOMPRESSED = 40 * 1024 * 1024


def read_pending_excel(content):
    if len(content) > MAX_FILE:
        raise ValidationError('Excel은 20MB 이하여야 합니다.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            if sum(i.file_size for i in archive.infolist()) > MAX_UNCOMPRESSED:
                raise ValidationError('압축 해제 크기가 너무 큽니다.')
        book = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, OSError, KeyError) as exc:
        raise ValidationError('정상적인 .xlsx 파일인지 확인해주세요.') from exc
    try:
        if 'PENDING' not in book.sheetnames:
            raise ValidationError('PENDING 시트가 필요합니다. 앱에서 연동용 Excel을 내려받아 사용해주세요.')
        sheet = book['PENDING']
        iterator = sheet.iter_rows()
        header_cells = next(iterator, ())
        headers = [str(c.value).strip() if c.value is not None else '' for c in header_cells]
        if len(headers) != len(set(headers)) or not set(FIELDS).issubset(headers):
            raise ValidationError('필수 열이 없거나 열 이름이 중복됩니다: ' + ', '.join(FIELDS))
        indexes = {name: headers.index(name) for name in FIELDS_WITH_VERSION if name in headers}
        rows = []; seen = set()
        for row_number, cells in enumerate(iterator, 2):
            if row_number > 5001:
                raise ValidationError('한 번에 5,000건 이하만 가져올 수 있습니다.')
            if all(c.value is None for c in cells):
                continue
            if any(c.data_type == 'f' for c in cells):
                raise ValidationError(f'{row_number}행에 수식이 있습니다. 연동 시트에는 값만 입력해주세요.')
            raw = {name: cells[index].value if index < len(cells) else None for name, index in indexes.items()}
            try:
                row = normalize_item(raw)
            except ValidationError as exc:
                raise ValidationError(f'{row_number}행: {exc}') from exc
            if row['ID'] in seen:
                raise ValidationError(f'{row_number}행: 중복 ID {row["ID"]}')
            seen.add(row['ID']); rows.append(row)
        return rows
    finally:
        book.close()


def pending_excel(rows):
    book = Workbook(); sheet = book.active; sheet.title = 'PENDING'
    sheet.append(FIELDS_WITH_VERSION)
    for row in rows:
        sheet.append([row.get(field, '') for field in FIELDS_WITH_VERSION])
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='CBD5E1'); cell.font = Font(bold=True)
    for column, width in zip('ABCDEFGH', [24, 18, 14, 14, 60, 20, 18, 12]):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'A2'; sheet.auto_filter.ref = sheet.dimensions
    guide = book.create_sheet('작성 안내')
    for line in [
        ['항목', '설명'], ['ID', '기존 ID는 유지하세요. 신규 항목은 고유 ID와 버전 0으로 입력합니다.'],
        ['버전', '기존 항목은 앱에서 내려받은 버전을 유지하세요. 직접 바꾸지 마세요.'],
        ['목표일', 'YYYY-MM-DD 또는 빈칸. 수식 대신 값만 입력합니다.'],
        ['분야', 'H/W, S/W, 공정, 전장, 기타'], ['상태', '미완료 또는 완료'],
        ['삭제', 'Excel에서 사라진 항목은 삭제 후보로만 안내하며 자동 삭제하지 않습니다.'],
        ['범위', '이 파일은 별도 Cloud 앱의 PENDING 연동 양식입니다. 기존 로컬 앱 원본 Excel과 다릅니다.']]:
        guide.append(line)
    guide.column_dimensions['A'].width = 18; guide.column_dimensions['B'].width = 100
    buf = BytesIO(); book.save(buf); return buf.getvalue()
