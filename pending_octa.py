"""Five-sheet Pending import; source values are retained alongside normalized fields."""
from datetime import date, datetime
from hashlib import sha256
from io import BytesIO
import re
from openpyxl import load_workbook
from domain import normalize_item, payload, ValidationError

SCHEMA_VERSION = 'pending-octa-1'
SHEETS = {'HW(설계)': 'H/W', '공정': '공정', '전장': '전장', 'SW': 'S/W', '기타': '기타'}


def original(value):
    return value.isoformat() if isinstance(value,(date,datetime)) else '' if value is None else str(value).strip()


def target_date(value):
    if value is None or str(value).strip() in ('','-'):
        return '', None
    try:
        if isinstance(value, datetime):
            parsed=value.date()
        elif isinstance(value,date):
            parsed=value
        else:
            match=re.fullmatch(r'(\d{2}|\d{4})[./-](\d{1,2})[./-](\d{1,2})',str(value).strip())
            if not match:raise ValueError()
            year,month,day=map(int,match.groups());year=2000+year if year<100 else year
            parsed=date(year,month,day)
        if not 2000 <= parsed.year <= 2100:raise ValueError()
        return parsed.isoformat(), None
    except ValueError:
        return '', '목표일을 날짜로 확정할 수 없어 기한 그래프에서 제외했습니다.'


def parse_pending(content, source_id='pending-excel'):
    book=load_workbook(BytesIO(content),read_only=True,data_only=True)
    try:
        if any(name not in book.sheetnames for name in SHEETS):
            raise ValidationError('PENDING 통합 양식의 HW(설계), 공정, 전장, SW, 기타 시트가 필요합니다.')
        incoming=[];details=[];warnings=[];seen=set()
        for name,category in SHEETS.items():
            sheet=book[name]
            if (sheet.max_row or 0)>5000 or (sheet.max_column or 0)>200:
                raise ValidationError(f'{name}: 허용 범위를 초과합니다.')
            rows=list(sheet.iter_rows(values_only=True))
            candidates=[i for i,row in enumerate(rows[:30]) if 'No.' in row and '호기' in row and '완료여부' in row]
            if len(candidates)!=1:raise ValidationError(f'{name}: 머리글을 찾을 수 없습니다.')
            header=candidates[0];names=[original(v) for v in rows[header]]
            required=['No.','호기','발생일','담당자','불합리 및 이슈 내용','목표일','완료여부']
            if any(names.count(field)!=1 for field in required):raise ValidationError(f'{name}: 필수 열이 없거나 중복됩니다.')
            indexes={field:names.index(field) for field in required}
            def get(row,field):
                i=indexes[field];return row[i] if i<len(row) else None
            for line,row in enumerate(rows[header+1:],header+2):
                number=get(row,'No.');description=get(row,'불합리 및 이슈 내용')
                if number is None and description is None:continue
                if not isinstance(number,(int,float)) or isinstance(number,bool) or int(number)!=number or number<1:
                    raise ValidationError(f'{name} {line}행: 고유 식별을 위한 No.가 필요합니다.')
                equipment=original(get(row,'호기'))
                occurred=original(get(row,'발생일'))
                normalized_occurrence,_=target_date(get(row,'발생일'))
                identity='|'.join([source_id,name,equipment,str(int(number)),normalized_occurrence or occurred])
                identifier='P-'+sha256(identity.encode()).hexdigest()[:32]
                if identifier in seen:raise ValidationError(f'{name} {line}행: 설비·No.·발생일이 같은 항목이 중복됩니다.')
                seen.add(identifier)
                completion=original(get(row,'완료여부')).upper()
                if completion not in ('O','X','완료','미완료'):raise ValidationError(f'{name} {line}행: 완료여부는 O/X 또는 완료/미완료여야 합니다.')
                target,warning=target_date(get(row,'목표일'))
                if warning:warnings.append({'시트':name,'행':line,'목표일 원문':original(get(row,'목표일')),'안내':warning})
                record=payload(normalize_item({'ID':identifier,'설비':equipment,'분야':category,'상태':'완료' if completion in ('O','완료') else '미완료','내용':original(description),'담당자':original(get(row,'담당자')),'목표일':target}))
                incoming.append(record)
                detail={'ID':identifier,'시트':name,'No.':int(number),'발생일':occurred,'목표일 원문':original(get(row,'목표일'))}
                for col,label in enumerate(names):
                    if label and col<len(row):detail[label]=original(row[col])
                details.append(detail)
        if not incoming:raise ValidationError('주요 5개 시트에 PENDING 항목이 없습니다.')
        if len(incoming)>5000:raise ValidationError('PENDING은 총 5,000건 이하여야 합니다.')
        return incoming,{'details':details,'warnings':warnings,'count':len(incoming),'done':sum(r['상태']=='완료' for r in incoming),'schema':SCHEMA_VERSION}
    finally:book.close()
