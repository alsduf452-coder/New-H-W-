from io import BytesIO
import hashlib
from openpyxl import Workbook, load_workbook
from fastapi.testclient import TestClient
from cloud_store import Repository
from domain import ValidationError
from pending_octa import parse_pending, SHEETS
from sync_api import create_app
from local_sync.agent import load_config, sync_once
from test_hw import hw_fixture
import pytest


def pending_fixture():
    book=Workbook();book.remove(book.active)
    for name in SHEETS:
        sheet=book.create_sheet(name)
        sheet.append(['No.','호기','발생일','담당자','불합리 및 이슈 내용','목표일','완료여부'])
        sheet.append([1,'DEMO-01','26.01.01','sample','sample issue','26.1.9','O'])
    book['SW'].append([1,'DEMO-01','26.01.02','sample','another issue','협의중','X'])
    out=BytesIO();book.save(out);return out.getvalue()


def test_pending_same_number_and_date_warnings():
    rows,info=parse_pending(pending_fixture())
    assert len(rows)==6 and info['done']==5
    assert len({r['ID'] for r in rows})==6
    assert len(info['warnings'])==1
    assert rows[0]['목표일']=='2026-01-09'
    assert next(r for r in rows if r['내용']=='another issue')['목표일']==''
    assert any(r['목표일 원문']=='협의중' for r in info['details'])


def test_ids_survive_sorting_and_content_edits_and_duplicate_keys_block():
    content=pending_fixture();before,_=parse_pending(content)
    book=load_workbook(BytesIO(content));sheet=book['SW']
    first=[c.value for c in sheet[2]];second=[c.value for c in sheet[3]]
    for i,v in enumerate(second,1):sheet.cell(2,i,v)
    for i,v in enumerate(first,1):sheet.cell(3,i,v)
    sheet.cell(3,5,'edited issue')
    out=BytesIO();book.save(out);after,_=parse_pending(out.getvalue())
    assert {r['ID'] for r in before}=={r['ID'] for r in after}
    sheet.append([c.value for c in sheet[3]])
    out=BytesIO();book.save(out)
    with pytest.raises(ValidationError,match='중복'):parse_pending(out.getvalue())


def test_two_sources_preserve_hw_and_pending_independently():
    repo=Repository();token='x'*43
    app=create_app(repo,token,mapping={'mode':'hw'},sources={'pending-excel':{'mode':'pending_octa'}})
    client=TestClient(app)
    def send(body,source,revision):
        return client.post('/v1/source/workbook',content=body,headers={'Authorization':'Bearer '+token,'If-Match':str(revision),'X-Sync-Source':source,'X-Content-SHA256':hashlib.sha256(body).hexdigest(),'Content-Type':'application/octet-stream'})
    assert send(hw_fixture(),'local-excel',0).status_code==200
    old=repo.local_previews('local-excel')
    assert send(pending_fixture(),'pending-excel',0).status_code==200
    assert len(repo.list_items())==6
    assert repo.local_previews('local-excel')==old
    assert repo.local_status('local-excel')['revision']==1
    assert repo.local_status('pending-excel')['pending_records']==6
    assert send(hw_fixture(),'local-excel',1).status_code==200
    assert len(repo.list_items())==6
    assert send(pending_fixture(),'unknown',0).status_code==404
    assert client.get('/v1/source/status',headers={'Authorization':'Bearer '+token}).json()['mode']=='hw'
    # PENDING conflict does not affect the H/W source.
    row=repo.list_items()[0];repo.save_item(dict(row,내용='app edit'),'editor',row['버전'])
    book=load_workbook(BytesIO(pending_fixture()))
    for sheet in book:
        sheet.cell(2,5,'excel edit')
    out=BytesIO();book.save(out)
    result=send(out.getvalue(),'pending-excel',1)
    assert result.json()['pending_status']=='conflict'
    assert repo.local_status('local-excel')['revision']==2


def test_optional_second_file_config(tmp_path):
    config=tmp_path/'config.toml'
    config.write_text(f"server_url='https://example.test'\nexcel_path='{tmp_path}/hw.xlsx'\npending_excel_path='{tmp_path}/pending.xlsx'\n")
    loaded=load_config(config)
    assert loaded['pending_source_id']=='pending-excel'
    assert str(loaded['pending_excel_path']).endswith('pending.xlsx')


def test_agent_routes_pending_header(tmp_path):
    repo=Repository();token='x'*43
    client=TestClient(create_app(repo,token,mapping={'mode':'hw'},sources={'pending-excel':{'mode':'pending_octa'}}))
    class Transport:
        def get(self,url,headers,**kwargs):return client.get('/v1/source/status',headers=headers)
        def post(self,url,headers,data,**kwargs):return client.post('/v1/source/workbook',headers=headers,content=data)
    path=tmp_path/'pending.xlsx';path.write_bytes(pending_fixture())
    config={'server_url':'https://example.test','excel_path':path,'source_id':'pending-excel'}
    assert '6건' in sync_once(config,token,Transport(),0)
    assert sync_once(config,token,Transport(),0)=='변경 없음'
    assert repo.local_status('local-excel')['revision']==0
