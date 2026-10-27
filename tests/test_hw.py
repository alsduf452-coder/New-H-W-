from io import BytesIO
import hashlib
from openpyxl import Workbook, load_workbook
from fastapi.testclient import TestClient
import pytest
from cloud_store import Repository
from domain import ValidationError
from hw_excel import parse_hw
from sync_api import create_app


def hw_fixture():
    book=Workbook();s=book.active;s.title='이력'
    headers=['Chamber','일자','변경이력','Diffuser','Backing Plate','Susceptor','Shadow Frame','RFG','Matcher','Reactor','Generator']
    for col,value in enumerate(headers,2): s.cell(3,col,value)
    for index,chamber in enumerate('ABCDE',4):
        s.cell(index,2,chamber);s.cell(index,3,'2000-01-01')
        for col in range(5,13):s.cell(index,col,f'SAMPLE-{chamber}-{col}')
    s['D4']='sample setup';s.merge_cells('D4:D8')
    s['B9']='D';s['C9']='2000-02-01';s['D9']='sample replacement';s['E9']='SAMPLE-REPLACED'
    for name in ['Spacing','Diff Flatness','Susc Flatness']:
        s=book.create_sheet(name);spacing=name=='Spacing'
        for index,chamber in enumerate('ABCDE'):
            c=(2 if spacing else 3)+9*index
            s.cell(3 if spacing else 4,c,'PC_'+chamber)
            anchors=[7] if spacing else [5,14]
            for r in anchors:
                s.cell(r,c,'Date');s.cell(r,c+1,'2000-01-01')
                s.cell(r+1,c,'Diff. S/N');s.cell(r+1,c+1,'SAMPLE-D')
                s.cell(r+1,c+3,'Susc. S/N');s.cell(r+1,c+4,'SAMPLE-S')
                if spacing:
                    s.cell(r-1,c,'sample spacing')
                    labels=[['①','⑥','②'],['⑨','⑤','⑦'],['④','⑧','③']]
                    for i in range(3):
                        for j in range(3):
                            s.cell(r+3+2*i,c+1+2*j,labels[i][j])
                            s.cell(r+4+2*i,c+1+2*j,10*i+j)
                else:
                    s.cell(r+2,c,'sample flatness')
                    for i in range(5):
                        for j in range(7):
                            if name=='Susc Flatness' and r==5 and j==3:continue
                            s.cell(r+3+i,c+j,-(10*i+j))
    out=BytesIO();book.save(out);return out.getvalue()


def test_hw_values_positions_missing_and_same_day_records():
    hw=parse_hw(hw_fixture())
    assert len(hw['history'])==6
    assert all(e['reason']=='sample setup' for e in hw['history'][:5])
    assert hw['installed']['D']['Diffuser']['sn']=='SAMPLE-REPLACED'
    assert hw['installed']['D']['Backing Plate']['sn']=='SAMPLE-D-6'
    flat=[r for r in hw['measurements'] if r['kind']=='Diffuser Flatness' and r['chamber']=='A']
    assert len(flat)==2 and flat[0]['id']!=flat[1]['id']
    assert flat[0]['grid'][0][0]==0
    assert flat[0]['center']==-23 and flat[0]['cells'][2][3]=='F10'
    susc=next(r for r in hw['measurements'] if r['kind']=='Susceptor Flatness')
    assert susc['center'] is None and susc['count']==30
    spacing=next(r for r in hw['measurements'] if r['kind']=='Spacing')
    assert spacing['grid']==[[0,1,2],[10,11,12],[20,21,22]]


def test_changed_layout_fails_closed():
    book=load_workbook(BytesIO(hw_fixture()));book['Spacing']['C10']='wrong'
    out=BytesIO();book.save(out)
    with pytest.raises(ValidationError,match='측정점 배치'):
        parse_hw(out.getvalue())


def test_hw_api_never_changes_pending_and_updates_snapshot():
    repo=Repository();token='x'*43
    api=TestClient(create_app(repo,token,mapping={'mode':'hw'}))
    content=hw_fixture()
    def send(content,revision):
        return api.post('/v1/source/workbook',content=content,headers={'Authorization':'Bearer '+token,'If-Match':str(revision),'Content-Type':'application/octet-stream','X-Content-SHA256':hashlib.sha256(content).hexdigest()})
    result=send(content,0)
    assert result.status_code==200
    assert result.json()['pending_status']=='not_applicable'
    assert result.json()['warning'] is None
    assert repo.list_items()==[]
    book=load_workbook(BytesIO(content));book['Diff Flatness']['F10']=-77
    out=BytesIO();book.save(out)
    assert send(out.getvalue(),1).status_code==200
    data=repo.local_previews('local-excel')['__hw_data__']
    assert next(m for m in data['measurements'] if m['id']=='Diff Flatness:C5')['center']==-77
    assert repo.list_items()==[]
    assert send(b'bad file',2).status_code==422
    assert repo.local_status('local-excel')['revision']==2


def test_hw_ui():
    from streamlit.testing.v1 import AppTest
    app=AppTest.from_string('''
from hw_ui import render_hw
import streamlit as st
render_hw(st.session_state.hw)
''')
    app.session_state['hw']=parse_hw(hw_fixture())
    app.run(timeout=30)
    assert not app.exception
    assert [t.label for t in app.tabs]==['장착 부품','작업·교체 이력','Spacing','Diffuser Flatness','Susceptor Flatness']
    assert any('Center (R3C4)' in table.value.columns for table in app.dataframe)
